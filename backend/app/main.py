import csv
import io
import logging
import os
import secrets
from datetime import datetime, timedelta, timezone
from typing import Literal

from fastapi import Depends, FastAPI, File, HTTPException, Query, Request, Response, UploadFile
from fastapi.exceptions import RequestValidationError
from fastapi.responses import JSONResponse, StreamingResponse
from openpyxl import Workbook
from pydantic import BaseModel, ConfigDict, Field
from sqlalchemy import func, or_, select, text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session
from starlette.middleware.trustedhost import TrustedHostMiddleware
from starlette.staticfiles import StaticFiles

from .crawl_api import router as crawl_router
from .db import get_db
from .importing import FIELDS, identity, parse_csv, safe_cell, website
from .models import (
    Audit,
    CrawlEntry,
    CrawlJob,
    ImportBatch,
    Location,
    LoginAttempt,
    Observation,
    Source,
    User,
    now,
)
from .security import create_session, current_user, hash_password, verify_password
from .settings import allowed_hosts, allowed_origins

app = FastAPI(title="LeadFinder", version="0.2.0", docs_url=None, redoc_url=None, openapi_url=None)
app.include_router(crawl_router)
app.add_middleware(TrustedHostMiddleware, allowed_hosts=allowed_hosts())
logger = logging.getLogger("leadfinder")
ALLOWED_ORIGINS = allowed_origins()


@app.middleware("http")
async def request_context(request, call_next):
    request.state.request_id = secrets.token_hex(8)
    if request.method not in ("GET", "HEAD", "OPTIONS"):
        origin = request.headers.get("origin")
        if origin and origin not in ALLOWED_ORIGINS:
            return JSONResponse(
                {"message": "Nguồn yêu cầu không hợp lệ.", "code": "origin_denied"}, status_code=403
            )
    try:
        response = await call_next(request)
    except Exception:
        logger.exception("Unhandled error request_id=%s", request.state.request_id)
        response = JSONResponse(
            {
                "code": "internal_error",
                "message": "Có lỗi xử lý. Hãy thử lại hoặc kiểm tra nhật ký.",
                "request_id": request.state.request_id,
            },
            status_code=500,
        )
    response.headers["X-Request-ID"] = request.state.request_id
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["Cache-Control"] = "no-store"
    return response


@app.exception_handler(HTTPException)
async def http_error(request, exc):
    return JSONResponse(
        {
            "code": f"http_{exc.status_code}",
            "message": exc.detail,
            "request_id": request.state.request_id,
        },
        status_code=exc.status_code,
    )


@app.exception_handler(RequestValidationError)
async def validation_error(request, exc):
    return JSONResponse(
        {
            "code": "invalid_input",
            "message": "Dữ liệu chưa hợp lệ. Hãy kiểm tra các trường và độ dài.",
            "request_id": request.state.request_id,
        },
        status_code=422,
    )


def audit(db, user, action, entity_id, changes):
    db.add(Audit(user_id=user.id, action=action, entity_id=entity_id, changes=changes))


def commit(db):
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(
            409, "Dữ liệu đã thay đổi hoặc trùng định danh. Hãy tải lại và thử lại."
        )


class Credentials(BaseModel):
    username: str = Field(min_length=3, max_length=100, pattern=r"^[a-zA-Z0-9_.-]+$")
    password: str = Field(min_length=12, max_length=128)


@app.get("/api/v1/health/live")
def live():
    return {"status": "ok", "version": "0.2.0"}


@app.get("/api/v1/health/ready")
def ready(db: Session = Depends(get_db)):
    try:
        db.execute(text("SELECT 1"))
        db.execute(select(User.id).limit(1))
    except Exception:
        raise HTTPException(503, "Cơ sở dữ liệu hoặc migration chưa sẵn sàng.")
    return {"status": "ready"}


@app.get("/api/v1/setup")
def setup_status(db: Session = Depends(get_db)):
    return {"required": db.get(User, 1) is None}


@app.post("/api/v1/setup", status_code=201)
def setup(body: Credentials, response: Response, db: Session = Depends(get_db)):
    if db.get(User, 1):
        raise HTTPException(409, "Ứng dụng đã được thiết lập.")
    user = User(id=1, username=body.username, password_hash=hash_password(body.password))
    db.add(user)
    commit(db)
    return {"username": user.username, "csrf": create_session(db, user.id, response)}


@app.post("/api/v1/auth/login")
def login(body: Credentials, request: Request, response: Response, db: Session = Depends(get_db)):
    key = request.client.host if request.client else "local"
    attempt = db.get(LoginAttempt, key)
    instant = datetime.now(timezone.utc)
    if not attempt:
        attempt = LoginAttempt(
            key=key, attempts=0, reset_at=(instant + timedelta(minutes=15)).isoformat()
        )
        db.add(attempt)
    if datetime.fromisoformat(attempt.reset_at) < instant:
        attempt.attempts = 0
        attempt.reset_at = (instant + timedelta(minutes=15)).isoformat()
    if attempt.attempts >= 10:
        raise HTTPException(429, "Đăng nhập sai nhiều lần. Vui lòng thử lại sau 15 phút.")
    user = db.scalar(select(User).where(User.username == body.username))
    if not user or not verify_password(body.password, user.password_hash):
        attempt.attempts += 1
        commit(db)
        raise HTTPException(401, "Tên đăng nhập hoặc mật khẩu chưa đúng.")
    attempt.attempts = 0
    return {"username": user.username, "csrf": create_session(db, user.id, response)}


@app.get("/api/v1/auth/me")
def me(request: Request, user=Depends(current_user)):
    return {"username": user.username, "csrf": request.state.auth_session.csrf}


@app.post("/api/v1/auth/logout")
def logout(
    request: Request, response: Response, user=Depends(current_user), db: Session = Depends(get_db)
):
    db.delete(request.state.auth_session)
    db.commit()
    response.delete_cookie("lf_session", path="/")
    return {"ok": True}


@app.get("/api/v1/overview")
def overview(user=Depends(current_user), db: Session = Depends(get_db)):
    def count(*conditions):
        return db.scalar(select(func.count()).select_from(Location).where(*conditions))

    return {
        "total": count(),
        "with_phone": count(Location.phone != ""),
        "with_website": count(Location.website != ""),
        "tax_pending": count(Location.tax_code != "", Location.tax_status != "confirmed"),
        "imports": db.scalar(
            select(func.count()).select_from(ImportBatch).where(ImportBatch.confirmed)
        ),
    }


def location_view(row, allowed_sources=None):
    result = {
        key: getattr(row, key)
        for key in [
            "id",
            *FIELDS,
            "tax_status",
            "status",
            "notes",
            "version",
            "created_at",
            "updated_at",
        ]
    }
    if allowed_sources is not None:
        for field in FIELDS:
            if row.provenance.get(field) not in allowed_sources:
                result[field] = ""
    return result


def display_sources(db):
    return set(db.scalars(select(Source.id).where(Source.allow_display)))


def filtered_locations(q="", city="", status="", has_phone=False):
    stmt = select(Location)
    if q:
        escaped = q.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
        stmt = stmt.where(
            or_(
                Location.name.ilike(f"%{escaped}%", escape="\\"),
                Location.address.ilike(f"%{escaped}%", escape="\\"),
                Location.phone.ilike(f"%{escaped}%", escape="\\"),
            )
        )
    if city:
        stmt = stmt.where(Location.city == city)
    if status:
        stmt = stmt.where(Location.status == status)
    if has_phone:
        stmt = stmt.where(Location.phone != "")
    return stmt


@app.get("/api/v1/locations")
def locations(
    q: str = Query("", max_length=150),
    city: str = "",
    status: str = "",
    has_phone: bool = False,
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    user=Depends(current_user),
    db: Session = Depends(get_db),
):
    stmt = filtered_locations(q, city, status, has_phone)
    total = db.scalar(select(func.count()).select_from(stmt.subquery()))
    rows = db.scalars(
        stmt.order_by(Location.updated_at.desc(), Location.id)
        .offset((page - 1) * page_size)
        .limit(page_size)
    ).all()
    allowed = display_sources(db)
    cities = db.scalars(
        select(Location.city)
        .where(Location.city != "", Location.provenance["city"].as_string().in_(allowed))
        .distinct()
        .order_by(Location.city)
    ).all()
    return {
        "items": [location_view(row, allowed) for row in rows],
        "total": total,
        "page": page,
        "page_size": page_size,
        "cities": cities,
    }


@app.get("/api/v1/locations/{location_id}")
def detail(location_id: str, user=Depends(current_user), db: Session = Depends(get_db)):
    row = db.get(Location, location_id)
    if not row:
        raise HTTPException(404, "Không tìm thấy địa điểm.")
    allowed = display_sources(db)
    observations = db.scalars(
        select(Observation)
        .where(Observation.location_id == row.id, Observation.source_id.in_(allowed))
        .order_by(Observation.created_at.desc())
    ).all()
    logs = db.scalars(
        select(Audit).where(Audit.entity_id == row.id).order_by(Audit.created_at.desc()).limit(50)
    ).all()
    return {
        **location_view(row, allowed),
        "observations": [
            {
                k: getattr(o, k)
                for k in ["id", "field", "value", "source_id", "source_ref", "created_at"]
            }
            for o in observations
        ],
        "history": [
            {"action": a.action, "changes": a.changes, "created_at": a.created_at} for a in logs
        ],
    }


class LocationUpdate(BaseModel):
    model_config = ConfigDict(extra="forbid")
    version: int
    name: str = Field(min_length=1, max_length=300)
    address: str = Field(max_length=500)
    city: str = Field(max_length=150)
    category: str = Field(max_length=100)
    phone: str = Field(max_length=150)
    website: str = Field(max_length=500)
    tax_code: str = Field(max_length=50)
    tax_status: Literal["unverified", "confirmed", "conflict"]
    status: Literal[
        "new", "contacted", "interested", "follow_up", "won", "not_suitable", "do_not_contact"
    ]
    notes: str = Field(max_length=5000)


@app.patch("/api/v1/locations/{location_id}")
def update(
    location_id: str,
    body: LocationUpdate,
    user=Depends(current_user),
    db: Session = Depends(get_db),
):
    row = db.scalar(select(Location).where(Location.id == location_id).with_for_update())
    if not row:
        raise HTTPException(404, "Không tìm thấy địa điểm.")
    if row.version != body.version:
        raise HTTPException(409, "Hồ sơ đã được sửa ở nơi khác. Hãy tải lại.")
    values = body.model_dump(exclude={"version"})
    values = {k: v.strip() for k, v in values.items()}
    try:
        values["website"] = website(values["website"])
    except ValueError as exc:
        raise HTTPException(422, str(exc))
    if not values["name"]:
        raise HTTPException(422, "Tên quán không được để trống.")
    if not values["address"] and not values["website"]:
        raise HTTPException(422, "Cần địa chỉ hoặc website để nhận diện địa điểm.")
    if values["tax_status"] == "confirmed" and not values["tax_code"]:
        raise HTTPException(422, "Chưa có MST để xác nhận.")
    if values["tax_code"] != row.tax_code:
        values["tax_status"] = "unverified"
    changes = {}
    provenance = dict(row.provenance)
    for key, value in values.items():
        old = getattr(row, key)
        if value != old:
            changes[key] = {"before": old, "after": value}
            setattr(row, key, value)
            if key in FIELDS:
                provenance[key] = "manual"
                db.add(
                    Observation(
                        location_id=row.id,
                        field=key,
                        value=value,
                        source_id="manual",
                        source_ref="Chỉnh sửa bởi " + user.username + " / " + now(),
                    )
                )
    row.provenance = provenance
    row.identity_key = identity(values)
    if changes:
        row.version += 1
        row.updated_at = now()
        audit(db, user, "update_location", row.id, changes)
    commit(db)
    return location_view(row)


@app.post("/api/v1/imports/preview", status_code=201)
async def preview(
    file: UploadFile = File(...), user=Depends(current_user), db: Session = Depends(get_db)
):
    data = await file.read(2 * 1024 * 1024 + 1)
    if len(data) > 2 * 1024 * 1024:
        raise HTTPException(413, "File tối đa 2 MB.")
    try:
        rows, mapping = parse_csv(data)
    except ValueError as exc:
        raise HTTPException(422, str(exc))
    batch = ImportBatch(filename=(file.filename or "import.csv")[:200], rows=rows)
    db.add(batch)
    db.commit()
    return {
        "id": batch.id,
        "rows": rows,
        "mapping": mapping,
        "valid": sum(not r["errors"] for r in rows),
        "invalid": sum(bool(r["errors"]) for r in rows),
    }


class ConfirmImport(BaseModel):
    rights_confirmed: bool


@app.post("/api/v1/imports/{batch_id}/confirm")
def confirm(
    batch_id: str, body: ConfirmImport, user=Depends(current_user), db: Session = Depends(get_db)
):
    if not body.rights_confirmed:
        raise HTTPException(422, "Cần xác nhận quyền sử dụng danh sách trước khi nhập.")
    batch = db.scalar(select(ImportBatch).where(ImportBatch.id == batch_id).with_for_update())
    if not batch:
        raise HTTPException(404, "Không tìm thấy lần nhập.")
    if batch.confirmed:
        return batch.result
    source = db.get(Source, "csv")
    if not source or not source.allow_store:
        raise HTTPException(403, "Nguồn CSV chưa cho phép lưu dữ liệu.")
    created, existing, conflicts = 0, 0, 0
    for item in batch.rows:
        if item["errors"]:
            continue
        values = item["data"]
        key = identity(values)
        row = db.scalar(select(Location).where(Location.identity_key == key))
        if not row:
            row = Location(
                identity_key=key,
                **{**values, "name": values["name"] or values["website"]},
                provenance={k: "csv" for k in FIELDS if values[k]},
            )
            db.add(row)
            db.flush()
            created += 1
        else:
            existing += 1
        for field, value in values.items():
            if not value:
                continue
            if getattr(row, field) != value:
                conflicts += 1
            ref = f"CSV {batch.id} / dòng {item['line']}"
            db.add(
                Observation(
                    location_id=row.id, field=field, value=value, source_id="csv", source_ref=ref
                )
            )
        audit(db, user, "import_observations", row.id, {"batch": batch.id, "line": item["line"]})
    batch.confirmed = True
    batch.result = {
        "created": created,
        "existing": existing,
        "conflicts": conflicts,
        "skipped": sum(bool(r["errors"]) for r in batch.rows),
    }
    audit(db, user, "confirm_import", batch.id, batch.result)
    commit(db)
    return batch.result


@app.get("/api/v1/sources")
def sources(user=Depends(current_user), db: Session = Depends(get_db)):
    return [
        {
            "id": s.id,
            "name": s.name,
            "allow_store": s.allow_store,
            "allow_display": s.allow_display,
            "allow_export": s.allow_export,
        }
        for s in db.scalars(select(Source))
    ]


@app.get("/api/v1/exports/download")
def export(
    format: Literal["csv", "xlsx"] = "xlsx",
    q: str = "",
    city: str = "",
    status: str = "",
    has_phone: bool = False,
    include_do_not_contact: bool = False,
    crawl_job_id: str = "",
    include_unverified_tax: bool = False,
    user=Depends(current_user),
    db: Session = Depends(get_db),
):
    stmt = filtered_locations(q, city, status, has_phone)
    if crawl_job_id:
        if not db.get(CrawlJob, crawl_job_id):
            raise HTTPException(404, "Không tìm thấy đợt thu thập.")
        stmt = stmt.where(
            Location.id.in_(
                select(CrawlEntry.location_id).where(
                    CrawlEntry.job_id == crawl_job_id, CrawlEntry.state == "done"
                )
            )
        )
    if not include_do_not_contact:
        stmt = stmt.where(Location.status != "do_not_contact")
    if db.scalar(select(func.count()).select_from(stmt.subquery())) > 10000:
        raise HTTPException(422, "Bản đầu xuất tối đa 10.000 địa điểm. Hãy thu hẹp bộ lọc.")
    rows = db.scalars(stmt.order_by(Location.name)).all()
    permitted = {s.id for s in db.scalars(select(Source).where(Source.allow_export))}
    headers = [
        "Mã địa điểm",
        "Tên quán",
        "Địa chỉ",
        "Tỉnh/thành",
        "Loại quán",
        "SĐT",
        "Website",
        "MST (xem trạng thái xác minh)" if include_unverified_tax else "MST đã xác nhận",
        "Trạng thái",
        "Ghi chú",
    ]
    fields = ["name", "address", "city", "category", "phone", "website", "tax_code"]
    if include_unverified_tax:
        headers.append("Trạng thái xác minh MST")
    values, evidence_rows, omitted = [], [], 0
    for row in rows:
        record = [row.id]
        for field in fields:
            value = getattr(row, field)
            if value and row.provenance.get(field) not in permitted:
                omitted += 1
                value = ""
            if field == "tax_code" and row.tax_status != "confirmed" and not include_unverified_tax:
                value = ""
            record.append(value)
        record.extend([row.status, row.notes])
        if include_unverified_tax:
            record.append(row.tax_status)
        values.append([safe_cell(v) for v in record])
    if rows:
        observations = db.scalars(
            select(Observation).where(
                Observation.location_id.in_([r.id for r in rows]),
                Observation.source_id.in_(permitted),
            )
        ).all()
        for obs in observations:
            evidence_rows.append(
                [
                    safe_cell(v)
                    for v in [
                        obs.location_id,
                        obs.field,
                        obs.value,
                        obs.source_id,
                        obs.source_ref,
                        obs.created_at,
                    ]
                ]
            )
    audit(
        db,
        user,
        "export",
        "locations",
        {
            "format": format,
            "rows": len(rows),
            "omitted_fields": omitted,
            "include_do_not_contact": include_do_not_contact,
        },
    )
    db.commit()
    if format == "csv":
        output = io.StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow(headers)
        writer.writerows(values)
        data = output.getvalue().encode("utf-8-sig")
        media = "text/csv; charset=utf-8"
    else:
        book = Workbook()
        sheet = book.active
        sheet.title = "Địa điểm"
        sheet.append(headers)
        for value in values:
            sheet.append(value)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(
                55, max(18, max(len(str(c.value or "")) for c in column) + 2)
            )
        evidence_sheet = book.create_sheet("Nguồn theo trường")
        evidence_sheet.append(
            ["Mã địa điểm", "Trường", "Giá trị", "Nguồn", "Tham chiếu", "Thời gian"]
        )
        for record in evidence_rows:
            evidence_sheet.append(record)
        evidence_sheet.freeze_panes = "A2"
        desc = book.create_sheet("Mô tả cột")
        desc.append(
            [
                "Lưu ý",
                "Có bao gồm MST ứng viên; kiểm tra cột trạng thái xác minh trước khi sử dụng."
                if include_unverified_tax
                else "MST chưa xác nhận được để trống ở sheet Địa điểm. Sheet nguồn có thể chứa giá trị ứng viên.",
            ]
        )
        desc.append(["Trường bị loại do quyền nguồn", omitted])
        desc.append(["Không liên hệ", "Mặc định loại hồ sơ do_not_contact."])
        output = io.BytesIO()
        book.save(output)
        data = output.getvalue()
        media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return StreamingResponse(
        io.BytesIO(data),
        media_type=media,
        headers={
            "Content-Disposition": f'attachment; filename="leadfinder.{format}"',
            "X-Omitted-Fields": str(omitted),
        },
    )


# The Render image bundles the compiled frontend. Local Compose keeps nginx.
if os.getenv("FRONTEND_DIST"):
    app.mount("/", StaticFiles(directory=os.environ["FRONTEND_DIST"], html=True), name="frontend")
