import csv
import io

from openpyxl import load_workbook

from app.db import SessionLocal
from app.importing import identity, normalized, safe_cell
from app.models import Source


def upload(client, rows):
    stream = io.StringIO()
    writer = csv.writer(stream)
    writer.writerow(["name", "address", "city", "category", "phone", "website", "tax_code"])
    writer.writerows(rows)
    response = client.post(
        "/api/v1/imports/preview",
        files={"file": ("test.csv", stream.getvalue().encode("utf-8-sig"), "text/csv")},
    )
    assert response.status_code == 201, response.text
    return response.json()


def confirm(client, batch):
    return client.post(f"/api/v1/imports/{batch['id']}/confirm", json={"rights_confirmed": True})


def payload(row):
    return {
        key: row[key]
        for key in [
            "version",
            "name",
            "address",
            "city",
            "category",
            "phone",
            "website",
            "tax_code",
            "tax_status",
            "status",
            "notes",
        ]
    }


def test_setup_and_auth_boundaries(client):
    assert client.get("/api/v1/locations").status_code == 401
    assert (
        client.post("/api/v1/setup", json={"username": "admin", "password": "short"}).status_code
        == 422
    )
    setup = client.post(
        "/api/v1/setup", json={"username": "admin", "password": "TestPasswordOnly123!"}
    )
    assert setup.status_code == 201
    assert "httponly" in setup.headers["set-cookie"].lower()
    assert (
        client.post(
            "/api/v1/setup", json={"username": "admin", "password": "TestPasswordOnly123!"}
        ).status_code
        == 409
    )
    assert client.post("/api/v1/auth/logout").status_code == 403
    assert (
        client.post(
            "/api/v1/auth/logout", headers={"X-CSRF-Token": setup.json()["csrf"]}
        ).status_code
        == 200
    )
    assert client.get("/api/v1/auth/me").status_code == 401
    assert (
        client.post(
            "/api/v1/auth/login", json={"username": "admin", "password": "wrong-but-long-password"}
        ).status_code
        == 401
    )


def test_csv_validation_and_reimport(auth):
    rows = [
        ["Quán An", "1 Đường A", "Đà Nẵng", "Quán ăn", "0900000000", "", "0012345678"],
        ["Quán An", "2 Đường A", "Đà Nẵng", "Quán ăn", "0900000000", "", "0012345678"],
        ["Thiếu địa chỉ", "", "", "", "", "", ""],
    ]
    batch = upload(auth, rows)
    assert (batch["valid"], batch["invalid"]) == (2, 1)
    first = confirm(auth, batch)
    assert first.status_code == 200, first.text
    assert first.json()["created"] == 2
    assert confirm(auth, batch).json() == first.json()
    assert confirm(auth, upload(auth, rows)).json()["existing"] == 2
    listing = auth.get("/api/v1/locations").json()
    assert listing["total"] == 2
    assert listing["items"][0]["tax_status"] == "unverified"
    assert (
        len(auth.get("/api/v1/locations/" + listing["items"][0]["id"]).json()["observations"]) >= 6
    )


def test_manual_changes_conflicts_and_version(auth):
    batch = upload(auth, [["Quán A", "1 Đường A", "Đà Nẵng", "", "0900000000", "", ""]])
    confirm(auth, batch)
    place = auth.get("/api/v1/locations").json()["items"][0]
    change = {**payload(place), "phone": "0911111111", "notes": "Đã kiểm tra"}
    response = auth.patch("/api/v1/locations/" + place["id"], json=change)
    assert response.status_code == 200, response.text
    assert auth.patch("/api/v1/locations/" + place["id"], json=change).status_code == 409
    result = confirm(
        auth, upload(auth, [["Quán A", "1 Đường A", "Đà Nẵng", "", "0922222222", "", ""]])
    ).json()
    assert result["conflicts"] == 1
    updated = auth.get("/api/v1/locations/" + place["id"]).json()
    assert updated["phone"] == "0911111111"
    assert {"0900000000", "0911111111", "0922222222"} <= {
        o["value"] for o in updated["observations"] if o["field"] == "phone"
    }


def test_export_permissions_tax_and_formula(auth):
    confirm(
        auth,
        upload(
            auth,
            [['=HYPERLINK("bad")', "1 Đường A", "Đà Nẵng", "", "+84900000000", "", "0012345678"]],
        ),
    )
    response = auth.get("/api/v1/exports/download?format=xlsx")
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ["Địa điểm", "Nguồn theo trường", "Mô tả cột"]
    assert book["Địa điểm"]["B2"].value.startswith("'=")
    assert not book["Địa điểm"]["H2"].value
    assert book["Địa điểm"]["F2"].data_type != "f"
    with SessionLocal() as db:
        db.get(Source, "csv").allow_export = False
        db.commit()
    restricted = auth.get("/api/v1/exports/download?format=csv")
    assert int(restricted.headers["X-Omitted-Fields"]) >= 4
    assert "HYPERLINK" not in restricted.content.decode("utf-8-sig")


def test_do_not_contact_filter(auth):
    confirm(auth, upload(auth, [["Quán A", "1 Đường A", "", "", "", "", ""]]))
    row = auth.get("/api/v1/locations").json()["items"][0]
    assert (
        auth.patch(
            "/api/v1/locations/" + row["id"], json={**payload(row), "status": "do_not_contact"}
        ).status_code
        == 200
    )
    assert (
        len(
            list(
                csv.reader(
                    io.StringIO(
                        auth.get("/api/v1/exports/download?format=csv").content.decode("utf-8-sig")
                    )
                )
            )
        )
        == 1
    )
    assert auth.get("/api/v1/locations?status=do_not_contact").json()["total"] == 1


def test_source_rights_origin_and_limits(auth):
    batch = upload(auth, [["Quán A", "1 Đường A", "", "", "", "", ""]])
    assert (
        auth.post(
            "/api/v1/imports/" + batch["id"] + "/confirm", json={"rights_confirmed": False}
        ).status_code
        == 422
    )
    assert (
        auth.post(
            "/api/v1/auth/logout", headers={"Origin": "https://untrusted.example"}
        ).status_code
        == 403
    )
    assert auth.get("/api/v1/locations?page_size=101").status_code == 422
    with SessionLocal() as db:
        db.get(Source, "csv").allow_store = False
        db.commit()
    assert confirm(auth, batch).status_code == 403


def test_unicode_and_formula_helpers():
    assert normalized("Đà Nẵng") == "da nang"
    assert safe_cell("\t=1+1").startswith("'")
    base = {"name": "Quán An", "address": "1 Đường A", "city": "Đà Nẵng", "website": ""}
    assert identity(base) != identity({**base, "city": "Hà Nội"})
    assert identity(base) != identity({**base, "name": "Quán Ân"})


def test_tax_change_invalidates_confirmation_and_display_policy(auth):
    confirm(auth, upload(auth, [["Quán A", "1 Đường A", "", "", "", "", "0012345678"]]))
    row = auth.get("/api/v1/locations").json()["items"][0]
    approved = auth.patch(
        "/api/v1/locations/" + row["id"], json={**payload(row), "tax_status": "confirmed"}
    ).json()
    assert approved["tax_status"] == "confirmed"
    changed = auth.patch(
        "/api/v1/locations/" + row["id"], json={**payload(approved), "tax_code": "0098765432"}
    ).json()
    assert changed["tax_status"] == "unverified"
    with SessionLocal() as db:
        db.get(Source, "csv").allow_display = False
        db.commit()
    hidden = auth.get("/api/v1/locations/" + row["id"]).json()
    assert hidden["name"] == ""
    assert all(obs["source_id"] != "csv" for obs in hidden["observations"])
