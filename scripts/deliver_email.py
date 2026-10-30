"""Send a generated LeadFinder workbook through SMTP."""

import argparse
import os
import smtplib
import sys
from datetime import datetime
from email.message import EmailMessage
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]


def mail_settings():
    required = ["SMTP_HOST", "SMTP_USERNAME", "SMTP_PASSWORD", "MAIL_TO"]
    missing = [name for name in required if not os.getenv(name, "").strip()]
    if missing:
        raise RuntimeError("Chưa cấu hình GitHub Actions Secrets: " + ", ".join(missing))
    recipients = [
        address.strip()
        for address in os.environ["MAIL_TO"].replace(";", ",").split(",")
        if address.strip()
    ]
    if not recipients:
        raise RuntimeError("MAIL_TO chưa có địa chỉ email nhận kết quả.")
    port = int(os.getenv("SMTP_PORT", "").strip() or "587")
    if not 1 <= port <= 65535:
        raise RuntimeError("SMTP_PORT phải nằm trong khoảng 1–65535.")
    return {
        "host": os.environ["SMTP_HOST"].strip(),
        "port": port,
        "username": os.environ["SMTP_USERNAME"].strip(),
        "password": os.environ["SMTP_PASSWORD"],
        "sender": os.getenv("MAIL_FROM", "").strip() or os.environ["SMTP_USERNAME"].strip(),
        "recipients": recipients,
    }


def summary_lines(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Tổng hợp"]
    rows = sheet.iter_rows(values_only=True)
    next(rows, None)
    lines = ["Kết quả thu thập tự động của LeadFinder:", ""]
    for row in rows:
        values = list(row) + [None] * 10
        if not values[0] and not values[1]:
            continue
        lines.append(
            f"- {values[0]} · {values[1]}: {values[3]}; "
            f"đã lưu {values[5] or 0}/{values[2] or 0}; tìm thấy {values[4] or 0}; "
            f"bỏ qua {values[6] or 0}; lỗi {values[7] or 0}; "
            f"cảnh báo website {values[8] or 0}."
        )
        if values[9]:
            lines.append(f"  {values[9]}")
    if len(lines) == 2:
        lines.append("Workbook không có dòng tổng hợp.")
    return lines


def send(path):
    settings = mail_settings()
    local_time = datetime.now(ZoneInfo("Asia/Ho_Chi_Minh"))
    message = EmailMessage()
    message["Subject"] = f"LeadFinder — kết quả thu thập {local_time:%d/%m/%Y}"
    message["From"] = settings["sender"]
    message["To"] = ", ".join(settings["recipients"])
    lines = summary_lines(path)
    attachment = path.read_bytes()
    if len(attachment) <= 18 * 1024 * 1024:
        message.set_content("\n".join(lines + ["", "File Excel tổng hợp được đính kèm."]))
        message.add_attachment(
            attachment,
            maintype="application",
            subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            filename=path.name,
        )
    else:
        run_url = os.getenv("GITHUB_SERVER_URL", "https://github.com") + "/"
        run_url += os.getenv("GITHUB_REPOSITORY", "") + "/actions/runs/"
        run_url += os.getenv("GITHUB_RUN_ID", "")
        message.set_content(
            "\n".join(lines + ["", "File quá lớn để đính kèm; tải trong lần chạy Actions:", run_url])
        )

    if settings["port"] == 465:
        with smtplib.SMTP_SSL(settings["host"], settings["port"], timeout=30) as client:
            client.login(settings["username"], settings["password"])
            client.send_message(message)
    else:
        with smtplib.SMTP(settings["host"], settings["port"], timeout=30) as client:
            client.ehlo()
            client.starttls()
            client.ehlo()
            client.login(settings["username"], settings["password"])
            client.send_message(message)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--input", type=Path, default=ROOT / "action-results" / "leadfinder-nightly.xlsx"
    )
    args = parser.parse_args()
    try:
        send(args.input)
        print("Đã gửi workbook qua email.", flush=True)
        return 0
    except Exception as exc:
        print(f"Gửi email thất bại: {exc}", file=sys.stderr, flush=True)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
