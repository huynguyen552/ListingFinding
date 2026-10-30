"""Append LeadFinder run results to a shared Google spreadsheet."""

import json
import os
import sys
import urllib.error
import urllib.parse
import urllib.request
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
API_ROOT = "https://sheets.googleapis.com/v4"
META_HEADERS = ["GitHub Run ID", "Thời gian thu thập"]
SHEET_HEADERS = {
    "Leads": [
        "Từ khóa",
        "Khu vực",
        "Mã địa điểm",
        "Tên quán",
        "Địa chỉ",
        "Tỉnh/thành",
        "Loại quán",
        "SĐT",
        "Website",
        "MST (xem trạng thái xác minh)",
        "Trạng thái",
        "Ghi chú",
        "Trạng thái xác minh MST",
    ],
    "Sources": ["Từ khóa", "Khu vực", "Mã địa điểm", "Trường", "Giá trị", "Nguồn", "Tham chiếu", "Thời gian"],
    "Notifications": ["Từ khóa", "Khu vực", "Tên địa điểm", "Trạng thái", "Chi tiết"],
}


class GoogleSheets:
    def __init__(self, spreadsheet_id, access_token):
        self.spreadsheet_id = urllib.parse.quote(spreadsheet_id, safe="")
        self.access_token = access_token

    def request(self, method, path, payload=None):
        body = None if payload is None else json.dumps(payload, ensure_ascii=False).encode("utf-8")
        headers = {
            "Accept": "application/json",
            "Authorization": f"Bearer {self.access_token}",
        }
        if body is not None:
            headers["Content-Type"] = "application/json; charset=utf-8"
        request = urllib.request.Request(
            API_ROOT + path, data=body, headers=headers, method=method
        )
        try:
            with urllib.request.urlopen(request, timeout=45) as response:
                raw = response.read()
        except urllib.error.HTTPError as exc:
            details = exc.read().decode("utf-8", errors="replace")[:1000]
            raise RuntimeError(f"Google Sheets API trả HTTP {exc.code}: {details}") from None
        if not raw:
            return {}
        return json.loads(raw.decode("utf-8"))

    def ensure_tabs(self, names):
        result = self.request(
            "GET",
            f"/spreadsheets/{self.spreadsheet_id}?fields=sheets.properties.title",
        )
        existing = {
            sheet["properties"]["title"] for sheet in result.get("sheets", [])
        }
        missing = [name for name in names if name not in existing]
        if missing:
            self.request(
                "POST",
                f"/spreadsheets/{self.spreadsheet_id}:batchUpdate",
                {"requests": [{"addSheet": {"properties": {"title": name}}} for name in missing]},
            )

    def values(self, method, sheet, cell_range, payload=None, query=""):
        a1_range = urllib.parse.quote(f"'{sheet}'!{cell_range}", safe="")
        operation = ":append" if method == "POST" else ""
        query_suffix = f"?{query}" if query else ""
        return self.request(
            method,
            f"/spreadsheets/{self.spreadsheet_id}/values/{a1_range}{operation}{query_suffix}",
            payload,
        )

    def append_table(self, title, headers, rows):
        existing = self.values("GET", title, "1:1").get("values", [])
        if not existing:
            self.values(
                "PUT",
                title,
                "A1",
                {"majorDimension": "ROWS", "values": [headers]},
                "valueInputOption=RAW",
            )
        elif existing[0] != headers:
            raise RuntimeError(
                f"Tiêu đề tab {title} không khớp LeadFinder; giữ nguyên dữ liệu hiện có."
            )
        if rows:
            self.values(
                "POST",
                title,
                "A1",
                {"majorDimension": "ROWS", "values": rows},
                "valueInputOption=RAW&insertDataOption=INSERT_ROWS",
            )


def cell_value(value):
    if value is None:
        return ""
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def read_table(workbook, title, fallback):
    sheet = workbook[title]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows or not any(value is not None for value in rows[0]):
        return fallback, []
    headers = [cell_value(value) for value in rows[0]]
    data = [
        [cell_value(value) for value in row]
        for row in rows[1:]
        if any(value is not None and value != "" for value in row)
    ]
    return headers, data


def publish(path):
    spreadsheet_id = os.getenv("GOOGLE_SHEET_ID", "").strip()
    access_token = os.getenv("GOOGLE_ACCESS_TOKEN", "").strip()
    if not spreadsheet_id or not access_token:
        raise RuntimeError("Thiếu GOOGLE_SHEET_ID hoặc access token Google Sheets.")

    workbook = load_workbook(path, read_only=True, data_only=True)
    collected_at = datetime.now(ZoneInfo("Asia/Ho_Chi_Minh")).isoformat(timespec="seconds")
    run_id = os.getenv("GITHUB_RUN_ID", "local")

    summary_headers, summary_rows = read_table(workbook, "Tổng hợp", [])
    leads_headers, leads_rows = read_table(workbook, "Địa điểm", SHEET_HEADERS["Leads"])
    sources_headers, sources_rows = read_table(
        workbook, "Nguồn theo trường", SHEET_HEADERS["Sources"]
    )
    notices_headers, notices_rows = read_table(
        workbook, "Thông báo", SHEET_HEADERS["Notifications"]
    )
    workbook.close()

    sheet = GoogleSheets(spreadsheet_id, access_token)
    sheet.ensure_tabs(["Runs", "Leads", "Sources", "Notifications"])
    sheet.append_table(
        "Runs",
        META_HEADERS + summary_headers,
        [[run_id, collected_at, *row] for row in summary_rows],
    )
    for tab, headers, rows in (
        ("Leads", leads_headers, leads_rows),
        ("Sources", sources_headers, sources_rows),
        ("Notifications", notices_headers, notices_rows),
    ):
        sheet.append_table(
            tab,
            META_HEADERS + headers,
            [[run_id, collected_at, *row] for row in rows],
        )
    print(
        "Đã ghi Google Sheets: "
        f"{len(summary_rows)} dòng tổng hợp, {len(leads_rows)} địa điểm, "
        f"{len(sources_rows)} nguồn, {len(notices_rows)} thông báo.",
        flush=True,
    )


def main():
    path = Path(os.getenv("LEADFINDER_WORKBOOK", ROOT / "action-results" / "leadfinder-nightly.xlsx"))
    try:
        publish(path)
        return 0
    except Exception as exc:
        print(f"Ghi Google Sheets thất bại: {exc}", file=sys.stderr, flush=True)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
